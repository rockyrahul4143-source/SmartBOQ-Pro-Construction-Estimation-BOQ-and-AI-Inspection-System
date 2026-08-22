"""
Report Generator Service
========================
Generates PDF, Excel, and CSV reports for projects, quantities, BOQs, and costs.
Uses ReportLab for PDF, openpyxl for Excel.
"""
from __future__ import annotations
import io
import csv
from datetime import datetime
from typing import Optional

# PDF
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph,
    Spacer, HRFlowable,
)
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT

# Excel
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ── Colours ───────────────────────────────────────────
PRIMARY   = colors.HexColor("#1E3A5F")   # dark navy
SECONDARY = colors.HexColor("#2E86AB")   # blue
ACCENT    = colors.HexColor("#F0A500")   # amber
LIGHT     = colors.HexColor("#EFF6FF")   # pale blue
WHITE     = colors.white
GREY      = colors.HexColor("#6B7280")


def _header_footer(canvas, doc, title: str, project_name: str):
    canvas.saveState()
    w, h = A4
    # Header bar
    canvas.setFillColor(PRIMARY)
    canvas.rect(0, h - 2.5*cm, w, 2.5*cm, fill=1, stroke=0)
    canvas.setFillColor(WHITE)
    canvas.setFont("Helvetica-Bold", 14)
    canvas.drawString(1.5*cm, h - 1.5*cm, "SmartBOQ Pro")
    canvas.setFont("Helvetica", 10)
    canvas.drawRightString(w - 1.5*cm, h - 1.5*cm, title)
    canvas.setFont("Helvetica", 8)
    canvas.drawString(1.5*cm, h - 2.2*cm, f"Project: {project_name}")
    canvas.drawRightString(w - 1.5*cm, h - 2.2*cm,
                           f"Generated: {datetime.utcnow().strftime('%d %b %Y %H:%M')} UTC")
    # Footer
    canvas.setFillColor(GREY)
    canvas.setFont("Helvetica", 8)
    canvas.drawString(1.5*cm, 0.8*cm, "SmartBOQ Pro — Confidential")
    canvas.drawRightString(w - 1.5*cm, 0.8*cm, f"Page {doc.page}")
    canvas.restoreState()


def _styles():
    s = getSampleStyleSheet()
    s.add(ParagraphStyle("SectionHead", parent=s["Heading2"],
                         textColor=PRIMARY, spaceBefore=12, spaceAfter=4))
    s.add(ParagraphStyle("TableHead", parent=s["Normal"],
                         textColor=WHITE, alignment=TA_CENTER, fontSize=9))
    s.add(ParagraphStyle("CellRight", parent=s["Normal"],
                         alignment=TA_RIGHT, fontSize=8))
    return s


# ─────────────────────────────────────────────────────
# BOQ Report — PDF
# ─────────────────────────────────────────────────────
def generate_boq_pdf(boq, project) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        topMargin=3*cm, bottomMargin=2*cm,
        leftMargin=1.5*cm, rightMargin=1.5*cm,
    )
    styles = _styles()
    story = []

    # Project info block
    info = [
        ["Project:", project.project_name, "BOQ No:", boq.boq_number],
        ["Client:",  project.client_name,  "Revision:", str(boq.revision)],
        ["Location:", project.location,    "Status:", boq.status.value.upper()],
        ["Date:", datetime.utcnow().strftime("%d %b %Y"), "Currency:", boq.currency],
    ]
    info_table = Table(info, colWidths=[3*cm, 7*cm, 3*cm, 5*cm])
    info_table.setStyle(TableStyle([
        ("FONTNAME",  (0,0), (-1,-1), "Helvetica"),
        ("FONTSIZE",  (0,0), (-1,-1), 9),
        ("FONTNAME",  (0,0), (0,-1), "Helvetica-Bold"),
        ("FONTNAME",  (2,0), (2,-1), "Helvetica-Bold"),
        ("TEXTCOLOR", (0,0), (0,-1), PRIMARY),
        ("TEXTCOLOR", (2,0), (2,-1), PRIMARY),
        ("BOTTOMPADDING", (0,0), (-1,-1), 4),
    ]))
    story.append(info_table)
    story.append(Spacer(1, 0.4*cm))
    story.append(HRFlowable(width="100%", thickness=1, color=PRIMARY))
    story.append(Spacer(1, 0.3*cm))

    # BOQ table header
    col_widths = [1.5*cm, 7*cm, 1.5*cm, 2.5*cm, 2.5*cm, 2.5*cm]
    headers = ["Item\nNo", "Description", "Unit", "Quantity", "Rate\n(INR)", "Amount\n(INR)"]
    data = [headers]

    for item in boq.items:
        if item.is_heading:
            data.append([item.item_no, item.description.upper(), "", "", "", ""])
        else:
            data.append([
                item.item_no,
                item.description,
                item.unit,
                f"{item.quantity:,.3f}",
                f"{item.rate:,.2f}",
                f"{item.amount:,.2f}",
            ])

    # Totals
    data.append(["", "", "", "", "Sub-Total:", f"{boq.subtotal or 0:,.2f}"])
    data.append(["", "", "", "", f"Overhead ({boq.overhead_pct}%):", f"{boq.overhead_amount or 0:,.2f}"])
    data.append(["", "", "", "", f"Profit ({boq.profit_pct}%):", f"{boq.profit_amount or 0:,.2f}"])
    data.append(["", "", "", "", f"Contingency ({boq.contingency_pct}%):", f"{boq.contingency_amount or 0:,.2f}"])
    data.append(["", "", "", "", "GRAND TOTAL:", f"{boq.grand_total or 0:,.2f}"])

    tbl = Table(data, colWidths=col_widths, repeatRows=1)
    # Build style commands
    style_cmds = [
        ("FONTNAME",    (0,0), (-1,0),  "Helvetica-Bold"),
        ("FONTSIZE",    (0,0), (-1,-1), 8),
        ("BACKGROUND",  (0,0), (-1,0),  PRIMARY),
        ("TEXTCOLOR",   (0,0), (-1,0),  WHITE),
        ("ALIGN",       (3,0), (-1,-1), "RIGHT"),
        ("ALIGN",       (0,0), (1,-1),  "LEFT"),
        ("GRID",        (0,0), (-1,-1), 0.25, colors.lightgrey),
        ("ROWBACKGROUNDS", (0,1), (-1,-5), [WHITE, LIGHT]),
        ("TOPPADDING",  (0,0), (-1,-1), 3),
        ("BOTTOMPADDING",(0,0),(-1,-1), 3),
    ]
    # Colour heading rows
    for i, item in enumerate(boq.items, 1):
        if item.is_heading:
            style_cmds.append(("BACKGROUND", (0,i), (-1,i), SECONDARY))
            style_cmds.append(("TEXTCOLOR",  (0,i), (-1,i), WHITE))
            style_cmds.append(("FONTNAME",   (0,i), (-1,i), "Helvetica-Bold"))
            style_cmds.append(("SPAN",       (1,i), (-1,i)))

    # Grand total row styling
    gt_row = len(data) - 1
    style_cmds += [
        ("BACKGROUND", (0, gt_row), (-1, gt_row), ACCENT),
        ("FONTNAME",   (0, gt_row), (-1, gt_row), "Helvetica-Bold"),
        ("TEXTCOLOR",  (0, gt_row), (-1, gt_row), PRIMARY),
    ]
    tbl.setStyle(TableStyle(style_cmds))
    story.append(tbl)

    doc.build(
        story,
        onFirstPage=lambda c, d: _header_footer(c, d, f"BILL OF QUANTITIES — {boq.boq_number}", project.project_name),
        onLaterPages=lambda c, d: _header_footer(c, d, f"BILL OF QUANTITIES — {boq.boq_number}", project.project_name),
    )
    return buf.getvalue()


# ─────────────────────────────────────────────────────
# BOQ Report — Excel
# ─────────────────────────────────────────────────────
def generate_boq_excel(boq, project) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "BOQ"

    hdr_font  = Font(bold=True, color="FFFFFF", size=10)
    hdr_fill  = PatternFill("solid", fgColor="1E3A5F")
    sec_fill  = PatternFill("solid", fgColor="2E86AB")
    tot_fill  = PatternFill("solid", fgColor="F0A500")
    alt_fill  = PatternFill("solid", fgColor="EFF6FF")
    bd = Side(style="thin", color="D1D5DB")
    border = Border(left=bd, right=bd, top=bd, bottom=bd)
    center = Alignment(horizontal="center", vertical="center")
    right  = Alignment(horizontal="right")

    # Title
    ws.merge_cells("A1:F1")
    ws["A1"] = f"BILL OF QUANTITIES — {boq.boq_number}"
    ws["A1"].font = Font(bold=True, size=14, color="1E3A5F")
    ws["A1"].alignment = center

    ws["A2"], ws["B2"] = "Project:", project.project_name
    ws["A3"], ws["B3"] = "Client:",  project.client_name
    ws["D2"], ws["E2"] = "Status:",  boq.status.value.upper()
    ws["D3"], ws["E3"] = "Currency:", boq.currency

    # Header row
    row = 5
    headers = ["Item No", "Description", "Unit", "Quantity", "Rate (INR)", "Amount (INR)"]
    col_widths = [10, 45, 10, 14, 18, 18]
    for c, (h, w) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = center
        cell.border = border
        ws.column_dimensions[get_column_letter(c)].width = w

    row += 1
    alt = False
    for item in boq.items:
        if item.is_heading:
            for c in range(1, 7):
                cell = ws.cell(row=row, column=c)
                cell.fill = sec_fill
                cell.font = Font(bold=True, color="FFFFFF", size=9)
                cell.border = border
            ws.cell(row=row, column=1, value=item.item_no)
            ws.cell(row=row, column=2, value=item.description.upper())
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        else:
            fill = alt_fill if alt else PatternFill()
            alt = not alt
            vals = [item.item_no, item.description, item.unit,
                    item.quantity, item.rate, item.amount]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=row, column=c, value=v)
                cell.fill = fill
                cell.border = border
                cell.font = Font(size=9)
                if c >= 4:
                    cell.number_format = "#,##0.00"
                    cell.alignment = right
        row += 1

    # Totals
    totals = [
        ("", "", "", "", "Sub-Total", boq.subtotal or 0),
        ("", "", "", "", f"Overhead ({boq.overhead_pct}%)", boq.overhead_amount or 0),
        ("", "", "", "", f"Profit ({boq.profit_pct}%)", boq.profit_amount or 0),
        ("", "", "", "", f"Contingency ({boq.contingency_pct}%)", boq.contingency_amount or 0),
        ("", "", "", "", "GRAND TOTAL", boq.grand_total or 0),
    ]
    for t_row in totals:
        is_grand = t_row[4] == "GRAND TOTAL"
        for c, v in enumerate(t_row, 1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.border = border
            if is_grand:
                cell.fill = tot_fill
                cell.font = Font(bold=True, size=10, color="1E3A5F")
            if c >= 5:
                cell.number_format = "#,##0.00"
                cell.alignment = right
        row += 1

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ─────────────────────────────────────────────────────
# BOQ Report — CSV
# ─────────────────────────────────────────────────────
def generate_boq_csv(boq, project) -> str:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["SmartBOQ Pro — Bill of Quantities"])
    writer.writerow(["Project", project.project_name, "BOQ No", boq.boq_number])
    writer.writerow(["Client", project.client_name, "Status", boq.status.value])
    writer.writerow([])
    writer.writerow(["Item No", "Description", "Unit", "Quantity", "Rate", "Amount"])
    for item in boq.items:
        writer.writerow([
            item.item_no, item.description, item.unit,
            item.quantity if not item.is_heading else "",
            item.rate if not item.is_heading else "",
            item.amount if not item.is_heading else "",
        ])
    writer.writerow([])
    writer.writerow(["", "", "", "", "Sub-Total", boq.subtotal or 0])
    writer.writerow(["", "", "", "", f"Overhead ({boq.overhead_pct}%)", boq.overhead_amount or 0])
    writer.writerow(["", "", "", "", f"Profit ({boq.profit_pct}%)", boq.profit_amount or 0])
    writer.writerow(["", "", "", "", f"Contingency ({boq.contingency_pct}%)", boq.contingency_amount or 0])
    writer.writerow(["", "", "", "", "GRAND TOTAL", boq.grand_total or 0])
    return output.getvalue()


# ─────────────────────────────────────────────────────
# Quantity Report — PDF
# ─────────────────────────────────────────────────────
def generate_quantity_pdf(estimates, project) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            topMargin=3*cm, bottomMargin=2*cm,
                            leftMargin=1.5*cm, rightMargin=1.5*cm)
    story = []
    col_widths = [3*cm, 6*cm, 2*cm, 2.5*cm, 2.5*cm, 2.5*cm, 2.5*cm]
    headers = ["Work Type", "Description", "Unit", "Quantity",
               "Cement\n(Bags)", "Sand\n(CFT)", "Steel\n(kg)"]
    data = [headers]
    for e in estimates:
        data.append([
            e.work_type.value.replace("_", " ").title(),
            f"Qty: {e.quantity:,.3f} {e.unit}",
            e.unit,
            f"{e.quantity:,.3f}",
            f"{e.cement_bags or 0:,.1f}",
            f"{e.sand_cft or 0:,.1f}",
            f"{e.steel_kg or 0:,.1f}",
        ])
    tbl = Table(data, colWidths=col_widths, repeatRows=1)
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), PRIMARY),
        ("TEXTCOLOR",  (0,0), (-1,0), WHITE),
        ("FONTNAME",   (0,0), (-1,0), "Helvetica-Bold"),
        ("FONTSIZE",   (0,0), (-1,-1), 8),
        ("GRID",       (0,0), (-1,-1), 0.25, colors.lightgrey),
        ("ROWBACKGROUNDS", (0,1), (-1,-1), [WHITE, LIGHT]),
        ("ALIGN",      (3,0), (-1,-1), "RIGHT"),
        ("TOPPADDING", (0,0), (-1,-1), 3),
        ("BOTTOMPADDING",(0,0),(-1,-1),3),
    ]))
    story.append(tbl)
    doc.build(
        story,
        onFirstPage=lambda c, d: _header_footer(c, d, "QUANTITY REPORT", project.project_name),
        onLaterPages=lambda c, d: _header_footer(c, d, "QUANTITY REPORT", project.project_name),
    )
    return buf.getvalue()


# ─────────────────────────────────────────────────────
# Measurement Book Report — PDF
# ─────────────────────────────────────────────────────
def generate_measurement_pdf(book, project) -> bytes:
    """
    Professional Measurement Book PDF.
    Columns: Ref | Description | Nos | L | W | H/D | Qty | Unit | Formula
    """
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            topMargin=3*cm, bottomMargin=2*cm,
                            leftMargin=1.2*cm, rightMargin=1.2*cm)
    story = []

    # Project info
    info_data = [
        ["Project:", project.project_name, "MB No:", book.mb_number],
        ["Location:", getattr(project, "location", "—"), "Date:", (book.date_of_measurement or datetime.utcnow()).strftime("%d %b %Y")],
        ["Checked By:", book.checked_by or "—", "Title:", book.title],
    ]
    info_tbl = Table(info_data, colWidths=[3*cm, 8*cm, 2.5*cm, 5*cm])
    info_tbl.setStyle(TableStyle([
        ("FONTSIZE",  (0,0), (-1,-1), 8),
        ("FONTNAME",  (0,0), (0,-1), "Helvetica-Bold"),
        ("FONTNAME",  (2,0), (2,-1), "Helvetica-Bold"),
        ("TEXTCOLOR", (0,0), (0,-1), PRIMARY),
        ("TEXTCOLOR", (2,0), (2,-1), PRIMARY),
        ("BOTTOMPADDING", (0,0), (-1,-1), 3),
    ]))
    story.append(info_tbl)
    story.append(Spacer(1, 0.3*cm))
    story.append(HRFlowable(width="100%", thickness=1, color=PRIMARY))
    story.append(Spacer(1, 0.2*cm))

    # Table
    col_widths = [1.2*cm, 5.5*cm, 1.2*cm, 1.8*cm, 1.8*cm, 1.8*cm, 1.8*cm, 1.2*cm, 4.5*cm]
    headers = ["Ref", "Description", "Nos", "L", "W", "H/D", "Qty", "Unit", "Formula"]
    data = [headers]

    for item in (book.items or []):
        if item.is_heading:
            data.append([item.item_ref or "", item.description.upper(),
                         "", "", "", "", "", "", ""])
        else:
            qty = item.quantity or 0
            qty_str = f"({abs(qty):.3f})" if item.is_deduction else f"{qty:.3f}"
            data.append([
                item.item_ref or "",
                ("  ↕ " if item.is_deduction else "  ") + item.description,
                f"{item.nos:.0f}" if item.nos and item.nos != 1 else "",
                f"{item.length:.3f}" if item.length else "",
                f"{item.width:.3f}"  if item.width  else "",
                f"{item.height:.3f}" if item.height  else "",
                qty_str,
                item.unit or "",
                item.formula_display or "",
            ])

    tbl = Table(data, colWidths=col_widths, repeatRows=1)
    style_cmds = [
        ("BACKGROUND",  (0,0), (-1,0), PRIMARY),
        ("TEXTCOLOR",   (0,0), (-1,0), WHITE),
        ("FONTNAME",    (0,0), (-1,0), "Helvetica-Bold"),
        ("FONTSIZE",    (0,0), (-1,-1), 7),
        ("GRID",        (0,0), (-1,-1), 0.25, colors.lightgrey),
        ("ROWBACKGROUNDS", (0,1), (-1,-1), [WHITE, LIGHT]),
        ("ALIGN",       (2,0), (-1,-1), "RIGHT"),
        ("TOPPADDING",  (0,0), (-1,-1), 2),
        ("BOTTOMPADDING",(0,0),(-1,-1), 2),
    ]
    for i, item in enumerate(book.items or [], 1):
        if item.is_heading:
            style_cmds += [
                ("BACKGROUND", (0,i), (-1,i), SECONDARY),
                ("TEXTCOLOR",  (0,i), (-1,i), WHITE),
                ("FONTNAME",   (0,i), (-1,i), "Helvetica-Bold"),
                ("SPAN",       (1,i), (-1,i)),
            ]
        elif item.is_deduction:
            style_cmds.append(("TEXTCOLOR", (0,i), (-1,i), colors.red))

    tbl.setStyle(TableStyle(style_cmds))
    story.append(tbl)

    # Quantity summary
    story.append(Spacer(1, 0.4*cm))
    story.append(HRFlowable(width="100%", thickness=0.5, color=GREY))
    story.append(Spacer(1, 0.2*cm))

    summary: dict = {}
    for item in (book.items or []):
        if not item.is_heading and item.quantity:
            key = f"{item.description} ({item.unit or ''})"
            summary[key] = summary.get(key, 0.0) + (item.quantity or 0.0)

    if summary:
        sum_data = [["Description", "Unit", "Total Quantity"]]
        for desc, qty in summary.items():
            parts = desc.rsplit("(", 1)
            unit = parts[1].rstrip(")") if len(parts) > 1 else ""
            sum_data.append([parts[0].strip(), unit, f"{qty:.3f}"])
        sum_tbl = Table(sum_data, colWidths=[9*cm, 2*cm, 3.5*cm])
        sum_tbl.setStyle(TableStyle([
            ("BACKGROUND", (0,0), (-1,0), ACCENT),
            ("FONTNAME",   (0,0), (-1,0), "Helvetica-Bold"),
            ("FONTSIZE",   (0,0), (-1,-1), 8),
            ("GRID",       (0,0), (-1,-1), 0.25, colors.lightgrey),
            ("ALIGN",      (2,0), (-1,-1), "RIGHT"),
        ]))
        story.append(Paragraph("Quantity Summary", ParagraphStyle("sh", fontSize=9, textColor=PRIMARY, fontName="Helvetica-Bold")))
        story.append(Spacer(1, 0.2*cm))
        story.append(sum_tbl)

    doc.build(
        story,
        onFirstPage=lambda c, d: _header_footer(c, d, f"MEASUREMENT BOOK — {book.mb_number}", project.project_name),
        onLaterPages=lambda c, d: _header_footer(c, d, f"MEASUREMENT BOOK — {book.mb_number}", project.project_name),
    )
    return buf.getvalue()


# ─────────────────────────────────────────────────────
# Rate Analysis Report — PDF
# ─────────────────────────────────────────────────────
def generate_rate_analysis_pdf(ra, project=None) -> bytes:
    """
    Professional Rate Analysis sheet.
    One page per item: Material / Labour / Equipment breakdown → Final Rate.
    """
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            topMargin=3*cm, bottomMargin=2*cm,
                            leftMargin=1.5*cm, rightMargin=1.5*cm)
    story = []

    project_name = project.project_name if project else "All Projects"
    story.append(Spacer(1, 0.3*cm))

    # Item header
    hdr_data = [
        [f"Rate Analysis: {ra.title}", f"Unit: {ra.unit}"],
        [f"SOR Code: {getattr(ra.sor_item, 'item_code', '—')}", f"Date: {datetime.utcnow().strftime('%d %b %Y')}"],
    ]
    hdr_tbl = Table(hdr_data, colWidths=[12*cm, 6*cm])
    hdr_tbl.setStyle(TableStyle([
        ("FONTNAME",  (0,0), (0,-1), "Helvetica-Bold"),
        ("FONTSIZE",  (0,0), (-1,-1), 9),
        ("TEXTCOLOR", (0,0), (-1,-1), PRIMARY),
        ("BOTTOMPADDING", (0,0), (-1,-1), 3),
    ]))
    story.append(hdr_tbl)
    story.append(Spacer(1, 0.3*cm))
    story.append(HRFlowable(width="100%", thickness=1, color=PRIMARY))
    story.append(Spacer(1, 0.2*cm))

    col_widths = [3*cm, 7*cm, 2*cm, 2.5*cm, 2.5*cm, 2.5*cm]
    headers = ["Type", "Description", "Unit", "Qty", "Rate (INR)", "Amount (INR)"]
    data = [headers]

    for ctype, label in [("material", "A. MATERIAL"), ("labour", "B. LABOUR"), ("equipment", "C. EQUIPMENT")]:
        comps = [c for c in (ra.components or []) if c.component_type == ctype]
        if not comps:
            continue
        data.append([label, "", "", "", "", ""])
        for c in comps:
            data.append(["", c.description, c.unit,
                         f"{c.quantity:.3f}", f"{c.rate:,.2f}", f"{c.amount:,.2f}"])

    # Totals section
    data.append(["", "", "", "", "Material Total:", f"{ra.material_total:,.2f}"])
    data.append(["", "", "", "", "Labour Total:",   f"{ra.labour_total:,.2f}"])
    data.append(["", "", "", "", "Equipment Total:", f"{ra.equipment_total:,.2f}"])
    data.append(["", "", "", "", "Direct Cost:",    f"{ra.direct_cost:,.2f}"])
    data.append(["", "", "", "", f"Wastage ({ra.wastage_pct}%):",   f"{ra.wastage_amount:,.2f}"])
    data.append(["", "", "", "", f"Overhead ({ra.overhead_pct}%):", f"{ra.overhead_amount:,.2f}"])
    data.append(["", "", "", "", f"Profit ({ra.profit_pct}%):",     f"{ra.profit_amount:,.2f}"])
    data.append(["", "", "", "", f"FINAL RATE / {ra.unit}:", f"{ra.final_rate:,.2f}"])

    tbl = Table(data, colWidths=col_widths, repeatRows=1)
    n = len(data)
    style_cmds = [
        ("BACKGROUND",  (0,0), (-1,0), PRIMARY),
        ("TEXTCOLOR",   (0,0), (-1,0), WHITE),
        ("FONTNAME",    (0,0), (-1,0), "Helvetica-Bold"),
        ("FONTSIZE",    (0,0), (-1,-1), 8),
        ("GRID",        (0,0), (-1,-1), 0.25, colors.lightgrey),
        ("ROWBACKGROUNDS", (0,1), (-1,-9), [WHITE, LIGHT]),
        ("ALIGN",       (3,0), (-1,-1), "RIGHT"),
        ("TOPPADDING",  (0,0), (-1,-1), 3),
        ("BOTTOMPADDING",(0,0),(-1,-1), 3),
        # Grand total row
        ("BACKGROUND",  (0, n-1), (-1, n-1), ACCENT),
        ("FONTNAME",    (0, n-1), (-1, n-1), "Helvetica-Bold"),
        ("TEXTCOLOR",   (0, n-1), (-1, n-1), PRIMARY),
        ("FONTSIZE",    (0, n-1), (-1, n-1), 10),
    ]
    # Section heading rows
    row_idx = 1
    for ctype, label in [("material", "A. MATERIAL"), ("labour", "B. LABOUR"), ("equipment", "C. EQUIPMENT")]:
        comps = [c for c in (ra.components or []) if c.component_type == ctype]
        if comps:
            style_cmds += [
                ("BACKGROUND", (0, row_idx), (-1, row_idx), SECONDARY),
                ("TEXTCOLOR",  (0, row_idx), (-1, row_idx), WHITE),
                ("FONTNAME",   (0, row_idx), (-1, row_idx), "Helvetica-Bold"),
                ("SPAN",       (0, row_idx), (4, row_idx)),
            ]
            row_idx += len(comps) + 1
        else:
            row_idx += 1

    tbl.setStyle(TableStyle(style_cmds))
    story.append(tbl)

    doc.build(
        story,
        onFirstPage=lambda c, d: _header_footer(c, d, "RATE ANALYSIS", project_name),
        onLaterPages=lambda c, d: _header_footer(c, d, "RATE ANALYSIS", project_name),
    )
    return buf.getvalue()


# ─────────────────────────────────────────────────────
# RA Bill Report — PDF
# ─────────────────────────────────────────────────────
def generate_ra_bill_pdf(bill, boq, project) -> bytes:
    """
    Running Account Bill PDF.
    Standard format: BOQ Qty | Prev Qty | Current Qty | Cum Qty | Balance | Rate | Amount
    """
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            topMargin=3*cm, bottomMargin=2*cm,
                            leftMargin=1*cm, rightMargin=1*cm)
    story = []

    # Header info
    info_data = [
        ["Project:", project.project_name, "Bill No:", bill.bill_number],
        ["Contractor:", bill.contractor_name or "—", "Status:", bill.status.upper()],
        ["Period From:", str(bill.period_from or "—"), "Period To:", str(bill.period_to or "—")],
    ]
    info_tbl = Table(info_data, colWidths=[2.5*cm, 8*cm, 2.5*cm, 5*cm])
    info_tbl.setStyle(TableStyle([
        ("FONTNAME",  (0,0), (0,-1), "Helvetica-Bold"),
        ("FONTNAME",  (2,0), (2,-1), "Helvetica-Bold"),
        ("FONTSIZE",  (0,0), (-1,-1), 8),
        ("TEXTCOLOR", (0,0), (0,-1), PRIMARY),
        ("TEXTCOLOR", (2,0), (2,-1), PRIMARY),
        ("BOTTOMPADDING", (0,0), (-1,-1), 3),
    ]))
    story.append(info_tbl)
    story.append(Spacer(1, 0.3*cm))
    story.append(HRFlowable(width="100%", thickness=1, color=PRIMARY))
    story.append(Spacer(1, 0.2*cm))

    col_widths = [1.2*cm, 5*cm, 1.2*cm, 1.5*cm, 2*cm, 1.8*cm, 2*cm, 1.6*cm, 1.8*cm, 2.2*cm]
    headers = ["Item\nNo", "Description", "Unit", "Rate", "BOQ\nQty",
               "Prev\nQty", "Current\nQty", "Cum\nQty", "Balance\nQty", "Current\nAmt"]
    data = [headers]

    for item in (bill.items or []):
        if item.is_heading:
            data.append([item.item_no, item.description.upper()] + [""] * 8)
        else:
            data.append([
                item.item_no, item.description, item.unit,
                f"{item.rate:,.2f}",
                f"{item.boq_quantity:.3f}",
                f"{item.previous_quantity:.3f}",
                f"{item.current_quantity:.3f}",
                f"{item.cumulative_quantity:.3f}",
                f"{item.balance_quantity:.3f}",
                f"{item.current_amount:,.2f}",
            ])

    # Bill totals
    data.append(["", "", "", "", "", "", "", "", "Current Bill:", f"{bill.current_amount:,.2f}"])
    data.append(["", "", "", "", "", "", "", "", "Previous Cum.:", f"{bill.previous_amount:,.2f}"])
    data.append(["", "", "", "", "", "", "", "", "Deductions:", f"{bill.deductions:,.2f}"])
    data.append(["", "", "", "", "", "", "", "", "NET PAYABLE:", f"{bill.net_payable:,.2f}"])

    tbl = Table(data, colWidths=col_widths, repeatRows=1)
    n = len(data)
    style_cmds = [
        ("BACKGROUND",  (0,0), (-1,0), PRIMARY),
        ("TEXTCOLOR",   (0,0), (-1,0), WHITE),
        ("FONTNAME",    (0,0), (-1,0), "Helvetica-Bold"),
        ("FONTSIZE",    (0,0), (-1,-1), 7),
        ("GRID",        (0,0), (-1,-1), 0.25, colors.lightgrey),
        ("ROWBACKGROUNDS", (0,1), (-1,-5), [WHITE, LIGHT]),
        ("ALIGN",       (3,0), (-1,-1), "RIGHT"),
        ("TOPPADDING",  (0,0), (-1,-1), 2),
        ("BOTTOMPADDING",(0,0),(-1,-1), 2),
        ("BACKGROUND",  (0, n-1), (-1, n-1), ACCENT),
        ("FONTNAME",    (0, n-1), (-1, n-1), "Helvetica-Bold"),
        ("TEXTCOLOR",   (0, n-1), (-1, n-1), PRIMARY),
    ]
    for i, item in enumerate(bill.items or [], 1):
        if item.is_heading:
            style_cmds += [
                ("BACKGROUND", (0,i), (-1,i), SECONDARY),
                ("TEXTCOLOR",  (0,i), (-1,i), WHITE),
                ("FONTNAME",   (0,i), (-1,i), "Helvetica-Bold"),
                ("SPAN",       (1,i), (-1,i)),
            ]

    tbl.setStyle(TableStyle(style_cmds))
    story.append(tbl)

    doc.build(
        story,
        onFirstPage=lambda c, d: _header_footer(c, d, f"RA BILL — {bill.bill_number}", project.project_name),
        onLaterPages=lambda c, d: _header_footer(c, d, f"RA BILL — {bill.bill_number}", project.project_name),
    )
    return buf.getvalue()


# ─────────────────────────────────────────────────────
# Abstract of Cost — PDF
# ─────────────────────────────────────────────────────
def generate_abstract_pdf(boq, project) -> bytes:
    """
    Abstract of Cost — category-wise summary with percentages.
    """
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            topMargin=3*cm, bottomMargin=2*cm,
                            leftMargin=1.5*cm, rightMargin=1.5*cm)
    story = []
    story.append(Spacer(1, 0.3*cm))

    # Aggregate amounts by work_type / section heading
    sections: dict = {}
    current_section = "General"
    for item in (boq.items or []):
        if item.is_heading:
            current_section = item.description
            if current_section not in sections:
                sections[current_section] = 0.0
        else:
            sections[current_section] = sections.get(current_section, 0.0) + (item.amount or 0.0)

    subtotal = boq.subtotal or sum(sections.values())
    col_widths = [1.5*cm, 10*cm, 3*cm, 3*cm]
    headers = ["S.No", "Description of Work", "Amount (INR)", "% of Total"]
    data = [headers]

    for i, (section, amount) in enumerate(sections.items(), 1):
        pct = (amount / subtotal * 100) if subtotal else 0
        data.append([str(i), section, f"{amount:,.2f}", f"{pct:.1f}%"])

    data.append(["", "Sub-Total", f"{subtotal:,.2f}", "100.0%"])
    data.append(["", f"Overhead ({boq.overhead_pct}%)",    f"{boq.overhead_amount or 0:,.2f}", ""])
    data.append(["", f"Profit ({boq.profit_pct}%)",        f"{boq.profit_amount or 0:,.2f}",  ""])
    data.append(["", f"Contingency ({boq.contingency_pct}%)", f"{boq.contingency_amount or 0:,.2f}", ""])
    data.append(["", "GRAND TOTAL", f"{boq.grand_total or 0:,.2f}", ""])

    n = len(data)
    tbl = Table(data, colWidths=col_widths, repeatRows=1)
    tbl.setStyle(TableStyle([
        ("BACKGROUND",  (0,0), (-1,0), PRIMARY),
        ("TEXTCOLOR",   (0,0), (-1,0), WHITE),
        ("FONTNAME",    (0,0), (-1,0), "Helvetica-Bold"),
        ("FONTSIZE",    (0,0), (-1,-1), 9),
        ("GRID",        (0,0), (-1,-1), 0.4, colors.lightgrey),
        ("ROWBACKGROUNDS", (0,1), (-1,-6), [WHITE, LIGHT]),
        ("ALIGN",       (2,0), (-1,-1), "RIGHT"),
        ("FONTNAME",    (0, n-6), (-1, n-6), "Helvetica-Bold"),
        ("BACKGROUND",  (0, n-1), (-1, n-1), ACCENT),
        ("FONTNAME",    (0, n-1), (-1, n-1), "Helvetica-Bold"),
        ("TEXTCOLOR",   (0, n-1), (-1, n-1), PRIMARY),
        ("FONTSIZE",    (0, n-1), (-1, n-1), 11),
        ("TOPPADDING",  (0,0), (-1,-1), 4),
        ("BOTTOMPADDING",(0,0),(-1,-1), 4),
    ]))
    story.append(tbl)

    doc.build(
        story,
        onFirstPage=lambda c, d: _header_footer(c, d, f"ABSTRACT OF COST — {boq.boq_number}", project.project_name),
        onLaterPages=lambda c, d: _header_footer(c, d, f"ABSTRACT OF COST — {boq.boq_number}", project.project_name),
    )
    return buf.getvalue()


# ─────────────────────────────────────────────────────
# RA Bill — Excel
# ─────────────────────────────────────────────────────
def generate_ra_bill_excel(bill, project) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = f"RA Bill {bill.bill_number}"

    hdr_font = Font(bold=True, color="FFFFFF", size=9)
    hdr_fill = PatternFill("solid", fgColor="1E3A5F")
    sec_fill = PatternFill("solid", fgColor="2E86AB")
    tot_fill = PatternFill("solid", fgColor="F0A500")
    alt_fill = PatternFill("solid", fgColor="EFF6FF")
    bd = Side(style="thin", color="D1D5DB")
    border = Border(left=bd, right=bd, top=bd, bottom=bd)
    right = Alignment(horizontal="right")

    ws.merge_cells("A1:J1")
    ws["A1"] = f"RUNNING ACCOUNT BILL — {bill.bill_number}"
    ws["A1"].font = Font(bold=True, size=13, color="1E3A5F")
    ws["A1"].alignment = Alignment(horizontal="center")

    ws["A2"], ws["B2"] = "Project:", project.project_name
    ws["A3"], ws["B3"] = "Contractor:", bill.contractor_name or "—"
    ws["F2"], ws["G2"] = "Status:", bill.status.upper()
    ws["F3"], ws["G3"] = "Net Payable:", bill.net_payable

    row = 5
    headers = ["Item No", "Description", "Unit", "Rate",
               "BOQ Qty", "Prev Qty", "Current Qty", "Cum Qty", "Balance", "Current Amt"]
    widths   = [10, 40, 8, 12, 12, 12, 12, 12, 12, 16]

    for c, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = hdr_font; cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center"); cell.border = border
        ws.column_dimensions[get_column_letter(c)].width = w

    row += 1; alt = False
    for item in (bill.items or []):
        if item.is_heading:
            for c in range(1, 11):
                cell = ws.cell(row=row, column=c)
                cell.fill = sec_fill; cell.border = border
                cell.font = Font(bold=True, color="FFFFFF", size=9)
            ws.cell(row=row, column=1, value=item.item_no)
            ws.cell(row=row, column=2, value=item.description.upper())
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=10)
        else:
            fill = alt_fill if alt else PatternFill()
            alt = not alt
            vals = [item.item_no, item.description, item.unit, item.rate,
                    item.boq_quantity, item.previous_quantity, item.current_quantity,
                    item.cumulative_quantity, item.balance_quantity, item.current_amount]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=row, column=c, value=v)
                cell.fill = fill; cell.border = border; cell.font = Font(size=9)
                if c >= 4:
                    cell.number_format = "#,##0.00"; cell.alignment = right
        row += 1

    # Totals
    for label, value in [("Current Bill", bill.current_amount),
                          ("Deductions", bill.deductions),
                          ("NET PAYABLE", bill.net_payable)]:
        is_net = label == "NET PAYABLE"
        for c in range(1, 11):
            cell = ws.cell(row=row, column=c); cell.border = border
            if is_net:
                cell.fill = tot_fill; cell.font = Font(bold=True, size=10, color="1E3A5F")
        ws.cell(row=row, column=9, value=label).font = Font(bold=True)
        cell_v = ws.cell(row=row, column=10, value=value)
        cell_v.number_format = "#,##0.00"; cell_v.alignment = right
        row += 1

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ─────────────────────────────────────────────────────
# Measurement Book — Excel
# ─────────────────────────────────────────────────────
def generate_measurement_excel(book, project) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = book.mb_number

    hdr_font = Font(bold=True, color="FFFFFF", size=9)
    hdr_fill = PatternFill("solid", fgColor="1E3A5F")
    sec_fill = PatternFill("solid", fgColor="2E86AB")
    alt_fill = PatternFill("solid", fgColor="EFF6FF")
    ded_font = Font(color="CC0000", size=9, italic=True)
    bd = Side(style="thin", color="D1D5DB")
    border = Border(left=bd, right=bd, top=bd, bottom=bd)
    right = Alignment(horizontal="right")

    ws.merge_cells("A1:I1")
    ws["A1"] = f"MEASUREMENT BOOK — {book.mb_number} — {book.title}"
    ws["A1"].font = Font(bold=True, size=13, color="1E3A5F")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws["A2"], ws["B2"] = "Project:", project.project_name
    ws["A3"], ws["B3"] = "Checked By:", book.checked_by or "—"

    row = 5
    headers = ["Ref", "Description", "Nos", "L", "W", "H/D", "Qty", "Unit", "Formula"]
    widths   = [8, 40, 8, 10, 10, 10, 10, 8, 35]
    for c, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = hdr_font; cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center"); cell.border = border
        ws.column_dimensions[get_column_letter(c)].width = w

    row += 1; alt = False
    for item in (book.items or []):
        if item.is_heading:
            for c in range(1, 10):
                cell = ws.cell(row=row, column=c)
                cell.fill = sec_fill; cell.border = border
                cell.font = Font(bold=True, color="FFFFFF", size=9)
            ws.cell(row=row, column=1, value=item.item_ref or "")
            ws.cell(row=row, column=2, value=item.description.upper())
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=9)
        else:
            fill = alt_fill if alt else PatternFill()
            alt = not alt
            qty = item.quantity or 0
            font = ded_font if item.is_deduction else Font(size=9)
            vals = [item.item_ref or "",
                    ("  ↕ " if item.is_deduction else "  ") + item.description,
                    item.nos if item.nos and item.nos != 1 else None,
                    item.length, item.width, item.height,
                    -abs(qty) if item.is_deduction else qty,
                    item.unit or "",
                    item.formula_display or ""]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=row, column=c, value=v)
                cell.fill = fill; cell.border = border; cell.font = font
                if c in (3,4,5,6,7):
                    cell.number_format = "#,##0.000"; cell.alignment = right
        row += 1

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
